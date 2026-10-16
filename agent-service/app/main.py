import hashlib
import json
import os
import uuid
from pathlib import Path
from fastapi import FastAPI, File, Header, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel, Field
from . import store
from .worker import DurableWorker

TOKEN=os.getenv("AGENT_ACCESS_TOKEN","")
ALLOW_INSECURE_LOCAL=os.getenv("AGENT_ALLOW_INSECURE_LOCAL","false").lower()=="true"

app=FastAPI(title="R8 Durable Research Agent",version="8.0.0")

class ProjectIn(BaseModel):
    project_id:str|None=None
    title:str=Field(min_length=1,max_length=200)
    base_revision:int=0

class TaskIn(BaseModel):
    project_id:str
    goal:str=Field(min_length=1,max_length=20000)
    source_ids:list[str]
    rules:str=""
    rule_version:str=""
    model_policy:dict
    idempotency_key:str|None=None

class ResumeIn(BaseModel):
    checkpoint_id:str
    last_event_seq:int=0

def auth(authorization: str|None):
    if not TOKEN:
        if ALLOW_INSECURE_LOCAL:return
        raise HTTPException(503,"Agent 服务未配置 AGENT_ACCESS_TOKEN")
    if authorization!="Bearer "+TOKEN:
        raise HTTPException(401,"未授权")

def parse_rules(rules):
    if not rules:return {}
    try:return json.loads(rules)
    except Exception:return {"rule_text":rules,"required_evidence":True}

def extract_text(path,mime):
    if mime.startswith("text/") or mime in {"application/json","application/xml","text/csv"}:
        return Path(path).read_text(encoding="utf-8",errors="ignore")
    if mime=="application/pdf":
        try:
            import fitz
            doc=fitz.open(path)
            return "\n\n".join(page.get_text() for page in doc)
        except Exception as exc:
            raise ValueError("PDF_TEXT_EXTRACTION_FAILED") from exc
    if mime=="application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            from docx import Document
            doc=Document(path)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as exc:
            raise ValueError("DOCX_TEXT_EXTRACTION_FAILED") from exc
    if mime in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","application/vnd.ms-excel"}:
        try:
            from openpyxl import load_workbook
            wb=load_workbook(path,read_only=True,data_only=True)
            blocks=[]
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals=[str(x) for x in row if x is not None]
                    if vals:blocks.append(" | ".join(vals))
            return "\n".join(blocks)
        except Exception as exc:
            raise ValueError("XLSX_TEXT_EXTRACTION_FAILED") from exc
    raise ValueError("UNSUPPORTED_MIME")

def snapshot(tid):
    t=store.task(tid)
    if not t:raise HTTPException(404,"TASK_NOT_FOUND")
    out=dict(t)
    out["error_codes"]=json.loads(out.get("error_codes") or "[]")
    out["model_policy"]=json.loads(out.get("model_policy") or "{}")
    out["trusted_memory"]=[{
        "id":m["id"],"kind":m["kind"],"trust":m["trust"],
        "source_ref":m["source_ref"],"content_hash":m["content_hash"],
        "value":json.loads(m["value_json"])
    } for m in store.trusted_memories(tid)]
    out["evidence"]=store.evidence(tid)
    if t.get("draft_path") and Path(t["draft_path"]).exists():
        out["draft"]=Path(t["draft_path"]).read_text(encoding="utf-8")
    if t.get("verified_report_path") and Path(t["verified_report_path"]).exists():
        out["verified_report"]=Path(t["verified_report_path"]).read_text(encoding="utf-8")
    return out

@app.on_event("startup")
def startup():
    if os.getenv("AGENT_AUTOSTART_WORKER","true").lower()=="true":
        import threading
        threading.Thread(target=DurableWorker().run_forever,name="r8-worker",daemon=True).start()

@app.get("/")
def home():
    p=Path(__file__).resolve().parent.parent/"web"/"index.html"
    if p.exists():return FileResponse(p)
    return HTMLResponse("<h1>R8 Durable Research Agent</h1>")

@app.get("/health")
def health():
    return {"ok":True,"service":"r8-durable-research-agent","version":"8.0.0"}

@app.post("/v1/sync/push")
def sync_push(body:ProjectIn,authorization: str|None=Header(default=None)):
    auth(authorization)
    try:return store.sync_project(body.project_id or str(uuid.uuid4()),body.title,body.base_revision)
    except ValueError as exc:
        if str(exc)=="REVISION_CONFLICT":raise HTTPException(409,"REVISION_CONFLICT")
        raise HTTPException(400,str(exc))

@app.get("/v1/projects/{project_id}")
def get_project(project_id:str,authorization: str|None=Header(default=None)):
    auth(authorization)
    p=store.project(project_id)
    if not p:raise HTTPException(404,"PROJECT_NOT_FOUND")
    p["sources"]=store.sources_for_project(project_id)
    return p

@app.post("/v1/projects/{project_id}/sources")
async def upload_source(project_id:str,file:UploadFile=File(...),authorization: str|None=Header(default=None),x_upload_consent:str|None=Header(default=None),x_content_sha256:str|None=Header(default=None)):
    auth(authorization)
    if x_upload_consent!="yes":raise HTTPException(403,"UPLOAD_CONSENT_REQUIRED")
    if not x_content_sha256 or len(x_content_sha256)!=64:raise HTTPException(400,"CONTENT_SHA256_REQUIRED")
    if not store.project(project_id):raise HTTPException(404,"PROJECT_NOT_FOUND")
    target=store.OBJECTS/x_content_sha256
    tmp=store.OBJECTS/(".incoming-"+str(uuid.uuid4()))
    size=0;h=hashlib.sha256()
    with tmp.open("wb") as out:
        while True:
            block=await file.read(1024*1024)
            if not block:break
            size+=len(block)
            if size>25*1024*1024:
                tmp.unlink(missing_ok=True);raise HTTPException(413,"SOURCE_TOO_LARGE")
            h.update(block);out.write(block)
    digest=h.hexdigest()
    if digest!=x_content_sha256:
        tmp.unlink(missing_ok=True);raise HTTPException(400,"CONTENT_SHA256_MISMATCH")
    if target.exists():
        old_hash,old_size=store.file_sha(target)
        if old_hash!=digest or old_size!=size:
            tmp.unlink(missing_ok=True);raise HTTPException(409,"EXISTING_SOURCE_CORRUPTED")
        tmp.unlink(missing_ok=True)
    else:tmp.replace(target)
    try:text=extract_text(target,file.content_type or "application/octet-stream")
    except ValueError as exc:
        target.unlink(missing_ok=True);raise HTTPException(415,str(exc))
    text_path=store.TEXTS/(digest+".txt")
    text_path.write_text(text,encoding="utf-8")
    sid=str(uuid.uuid4())
    try:row=store.add_source(project_id,file.filename or "source",file.content_type or "application/octet-stream",digest,size,target,text_path,sid)
    except ValueError as exc:raise HTTPException(409,str(exc))
    return {"id":row["id"],"project_id":project_id,"sha256":digest,"size_bytes":size,"text_chars":len(text)}

@app.post("/v1/tasks")
def create_task(body:TaskIn,authorization: str|None=Header(default=None)):
    auth(authorization)
    policy=body.model_policy or {}
    if policy.get("mode")!="STRICT":raise HTTPException(400,"ONLY_STRICT_MODE_SUPPORTED")
    if policy.get("allow_model_fallback",False):raise HTTPException(400,"MODEL_FALLBACK_MUST_BE_DISABLED")
    if not policy.get("remote_data_consent",False):raise HTTPException(403,"REMOTE_DATA_CONSENT_REQUIRED")
    if not body.source_ids:raise HTTPException(400,"SOURCE_REQUIRED")
    try:
        t=store.create_task(body.project_id,body.goal,body.source_ids,body.rules,body.rule_version,
                            str(policy.get("model") or ""),policy,body.idempotency_key)
    except ValueError as exc:raise HTTPException(400,str(exc))
    return snapshot(t["id"])

@app.get("/v1/tasks/{task_id}")
def get_task(task_id:str,authorization: str|None=Header(default=None)):
    auth(authorization);return snapshot(task_id)

@app.get("/v1/tasks/{task_id}/events")
def get_events(task_id:str,after_seq:int=0,authorization: str|None=Header(default=None)):
    auth(authorization)
    if not store.task(task_id):raise HTTPException(404,"TASK_NOT_FOUND")
    return {"events":store.events(task_id,max(0,after_seq))}

@app.post("/v1/tasks/{task_id}/resume")
def resume(task_id:str,body:ResumeIn,authorization: str|None=Header(default=None)):
    auth(authorization);t=store.task(task_id)
    if not t:raise HTTPException(404,"TASK_NOT_FOUND")
    if body.checkpoint_id!=t.get("checkpoint_id"):raise HTTPException(409,"CHECKPOINT_MISMATCH")
    if int(body.last_event_seq)>int(t["last_event_seq"]):raise HTTPException(409,"EVENT_CURSOR_AHEAD")
    store.set_state(task_id,status="CHECKPOINTED",lease_owner=None,lease_until=None)
    store.append_event(task_id,"RESUME_ACCEPTED",{"checkpoint_id":body.checkpoint_id,"last_event_seq":body.last_event_seq})
    return snapshot(task_id)

@app.post("/v1/tasks/{task_id}/validate")
def validate(task_id:str,authorization: str|None=Header(default=None)):
    auth(authorization);t=store.task(task_id)
    if not t:raise HTTPException(404,"TASK_NOT_FOUND")
    from .validator import final_gate
    ok,reason=final_gate(t)
    if not ok:
        store.set_state(task_id,status="NEEDS_REVIEW")
        store.append_event(task_id,"VALIDATION_BLOCKED",{"reason":reason})
    return {"status":"PASS" if ok else "UNKNOWN","reason":reason}

